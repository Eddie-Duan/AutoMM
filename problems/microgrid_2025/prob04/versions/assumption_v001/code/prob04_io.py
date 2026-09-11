# -*- coding: utf-8 -*-
"""prob04 输入读取与 ``result4-2.xlsx`` / ``result4-3.xlsx`` 填报（两链共用）。

只读 ``data/附件2.xlsx``（实际负载/光伏）、``data/附件3.xlsx``（整点预报，仅 4-3）、
``data/附件4.xlsx``（实时电价，两链唯一价格来源）与模板 ``data/附件5/result4-2.xlsx`` /
``data/附件5/result4-3.xlsx``；不修改任何原始数据。

口径来源：prob04 ``assumption_v001``（AS01–AS23 + 团队 ``A8-(b)`` 六条派生）与
``formulations/formulation_v001``（§1.1–§1.5、§8.1–§8.3）。

纪律（AS02 / §8.7 黑名单）：
  * **附件 1 不使用**（其电价被附件 4 取代，其负载/光伏预测列属 prob01 语义）；
  * 决策层价格一律来自「历史实际价 + 预测器」（``\\hat p``，见 ``prob04_predict``），**不是**当天实际价；
    本模块只负责读取与填报，不构造任何决策价；
  * 附件 3 的 ``日期`` 列只在每日 ``0:00`` 行出现（1095/1460 行为空），实现按**前向填充**解析，
    并把 ``2025-1-1`` 规范为 ``2025-01-01``；
  * 两个交付工作簿**必须新写**入独立 ``output_directory``；模板只读打开。
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

LOAD_SHEET = "小区负载"
PV_SHEET = "光伏发电实际功率"
PLAN_SHEET = "计划购电量"
ADJUST_SHEET = "调整购电量"
STORAGE_SHEET = "充放电量"
EMERGENCY_SHEET = "紧急购电量"
PLAN_HEADER_FIRST = "日期\\时间"
PLAN_HEADER_ENERGY = "全天购电量"
PLAN_HEADER_COST = "全天购电费"
STORAGE_HEADER = ("日期", "时间段", "充电量", "放电量", "时刻", "储电量")
EMERGENCY_HEADER = ("日期", "购电时间段", "购电量")
BLOCK_ROWS = 24
BLOCKS_PER_DAY = 6
PERIODS_PER_DAY = 144
DAYS_FULL = 365
ZERO_INTERVAL_TEXT = "—"
POWER_DATE_HEADER = "日期\\时间"
PRICE_SHEET = "Sheet1"
DECISION_HOURS: tuple[int, ...] = (0, 6, 12, 18)
FC_HORIZON = 24
# 交付数值的有效位（承 prob03 勘误 R10）：求解器解带 1e-10 级噪声，交付一律按 6 位小数写入。
DELIVERY_DECIMALS = 6
STATE_MIN_KWH = 1200.0
STATE_MAX_KWH = 10800.0


class InputValidationError(RuntimeError):
    """输入结构与 accepted 假设不一致时严格失败。"""


@dataclass(frozen=True)
class Attachment2:
    """附件 2 的逐日实际负载/光伏功率（365 天 × 144 个 10 分钟点，单位 kW）。"""

    load_kw: np.ndarray
    pv_kw: np.ndarray
    dates: list[str]
    time_labels: list[str]
    md5: str

    @property
    def days(self) -> int:
        return int(self.load_kw.shape[0])

    @property
    def periods_per_day(self) -> int:
        return int(self.load_kw.shape[1])


@dataclass(frozen=True)
class Attachment3:
    """附件 3 的整点预报表 ``A[d, s, k]``（kW；``s`` 为 ``DECISION_HOURS`` 下标，``k = 1..24``）。"""

    fc_kw: np.ndarray          # (days, 4, 24)
    dates: list[str]
    slots: list[str]
    md5: str

    @property
    def days(self) -> int:
        return int(self.fc_kw.shape[0])

    def row(self, day: int, hour: int) -> np.ndarray:
        """返回第 ``day`` 天、发布时刻 ``hour`` 的 24 维整点预报向量。"""
        if hour not in DECISION_HOURS:
            raise ValueError(f"未知发布时刻：{hour}")
        return self.fc_kw[day, DECISION_HOURS.index(hour), :]


@dataclass(frozen=True)
class Attachment4:
    """附件 4 的逐日逐时段**实际**实时电价（365 天 × 144 个 10 分钟点，元/kWh）。"""

    price: np.ndarray
    dates: list[str]
    time_labels: list[str]
    header: list[str]
    md5: str

    @property
    def days(self) -> int:
        return int(self.price.shape[0])

    @property
    def periods_per_day(self) -> int:
        return int(self.price.shape[1])


def md5_of(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clock_label(value: Any) -> str:
    """把附件时间戳格式化为 ``H:MM``（小时不补零，与 AS01 左端点口径/模板列标签风格一致）。

    ``openpyxl`` 把附件 2/附件 4 的表头时刻读成 ``datetime.time``，其 ``str()`` 会带秒
    （``datetime.time(10, 0)`` → ``"10:00:00"``），与 ``AS01`` 的 ``H:MM``（``"10:00"``）不同；
    本函数统一为 ``H:MM``（字符串原样 strip 返回，保留 ``"0:00+1"`` 一类模板写法）。
    """
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return f"{int(value.hour)}:{int(value.minute):02d}"
    return str(value).strip()


def _date_label(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def canonical_date(value: Any) -> str:
    """把附件 3 的日期（``2025-1-1`` 字符串或 datetime）规范为 ``YYYY-MM-DD``；空白返回空串。"""
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text:
        return ""
    parts = text.split("-") if "-" in text else text.split("/")
    if len(parts) == 3:
        try:
            return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
        except ValueError:
            return text
    return text


def _slot_hour(value: Any) -> int:
    """把附件 3 的 ``预报时刻`` 解析为整点小时数（``0:00`` → 0）。"""
    if value is None:
        raise InputValidationError("附件 3 存在空的预报时刻")
    if hasattr(value, "strftime"):
        return int(value.strftime("%H"))
    text = str(value).strip()
    if not text:
        raise InputValidationError("附件 3 存在空的预报时刻")
    head = text.split(":")[0].strip()
    try:
        return int(head)
    except ValueError as exc:
        raise InputValidationError(f"附件 3 预报时刻无法解析：{value!r}") from exc


def _validate_dates(dates: Sequence[str], *, source: str) -> None:
    try:
        parsed = [datetime.strptime(item, "%Y-%m-%d").date() for item in dates]
    except ValueError as exc:
        raise InputValidationError(f"{source}日期列无法解析：{exc}") from exc
    for index in range(1, len(parsed)):
        if (parsed[index] - parsed[index - 1]).days != 1:
            raise InputValidationError(f"{source}日期不连续：{dates[index - 1]} -> {dates[index]}")
    if parsed[0].isoformat() != "2025-01-01":
        raise InputValidationError(f"{source}首日不是 2025-01-01：{dates[0]}")


def _read_power_sheet(
    sheet: Any, *, name: str, expected_days: int, expected_periods: int
) -> tuple[list[str], list[str], np.ndarray]:
    rows = [row for row in sheet.iter_rows(min_row=1, max_col=expected_periods + 1, values_only=True)]
    if len(rows) != expected_days + 1:
        raise InputValidationError(f"附件 2『{name}』行数 {len(rows)} != 表头 + {expected_days}")
    # 表头时刻列经 ``_clock_label`` 规范为 ``H:MM``（openpyxl 的 ``datetime.time`` 直接 str 会带秒，
    # 导致表 1 的「attachment_timestamp 为区间左端点」硬检查失败；2026-09-11 修复）。
    header = ["" if item is None else _clock_label(item) for item in rows[0]]
    if header[0].strip() != POWER_DATE_HEADER:
        raise InputValidationError(f"附件 2『{name}』首列表头不是 {POWER_DATE_HEADER}：{header[0]!r}")
    if len(header) != expected_periods + 1:
        raise InputValidationError(f"附件 2『{name}』列数 {len(header)} != {expected_periods + 1}")
    body = rows[1:]
    dates = [_date_label(row[0]) for row in body]
    try:
        values = np.array([[float(item) for item in row[1:]] for row in body], dtype=float)
    except (TypeError, ValueError) as exc:
        raise InputValidationError(f"附件 2『{name}』存在非数值单元格：{exc}") from exc
    if values.shape != (expected_days, expected_periods):
        raise InputValidationError(f"附件 2『{name}』矩阵形状 {values.shape} 不符")
    if not np.all(np.isfinite(values)):
        raise InputValidationError(f"附件 2『{name}』含 NaN/Inf")
    if np.any(values < 0):
        raise InputValidationError(f"附件 2『{name}』存在负功率")
    return header, dates, values


def read_attachment2(
    path: Path, expected_days: int = DAYS_FULL, expected_periods: int = PERIODS_PER_DAY
) -> Attachment2:
    """读取附件 2 的『小区负载』与『光伏发电实际功率』（均为实际值）。"""
    if not path.is_file():
        raise InputValidationError(f"附件 2 不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        titles = list(workbook.sheetnames)
        for name in (LOAD_SHEET, PV_SHEET):
            if name not in titles:
                raise InputValidationError(f"附件 2 缺少工作表『{name}』：{titles}")
        load_header, load_dates, load = _read_power_sheet(
            workbook[LOAD_SHEET], name=LOAD_SHEET, expected_days=expected_days, expected_periods=expected_periods
        )
        pv_header, pv_dates, pv = _read_power_sheet(
            workbook[PV_SHEET], name=PV_SHEET, expected_days=expected_days, expected_periods=expected_periods
        )
    finally:
        workbook.close()
    if load_dates != pv_dates:
        raise InputValidationError("附件 2 两个工作表的日期序列不一致")
    _validate_dates(load_dates, source="附件 2 ")
    if load_header[1:] != pv_header[1:]:
        raise InputValidationError("附件 2 两个工作表的时间标签不一致")
    return Attachment2(
        load_kw=load,
        pv_kw=pv,
        dates=load_dates,
        time_labels=[str(item) for item in load_header[1:]],
        md5=md5_of(path),
    )


def read_attachment3(
    path: Path, *, expected_days: int = DAYS_FULL, hours: Sequence[int] = DECISION_HOURS
) -> Attachment3:
    """读取附件 3 的整点预报表，返回 ``(days, 4, 24)`` 张量与规范日期/时刻序列。

    ``日期`` 列只在每日首个发布时刻行出现（其余为空白），实现按**前向填充**解析（团队勘误 prob03 E1）；
    每日必须恰有 ``len(hours)`` 行且发布时刻按 ``0:00 → 6:00 → 12:00 → 18:00`` 升序。
    """
    if not path.is_file():
        raise InputValidationError(f"附件 3 不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [row for row in sheet.iter_rows(min_row=1, max_col=2 + FC_HORIZON, values_only=True)]
    finally:
        workbook.close()
    if len(rows) < 2:
        raise InputValidationError("附件 3 没有数据行")
    header = [str(item).strip() if item is not None else "" for item in rows[0]]
    if header[:2] != ["日期", "预报时刻"]:
        raise InputValidationError(f"附件 3 表头前两列异常：{header[:2]}")
    body = rows[1:]
    expected_rows = expected_days * len(hours)
    if len(body) != expected_rows:
        raise InputValidationError(f"附件 3 数据行数 {len(body)} != 期望 {expected_rows}")

    current = ""
    filled: list[str] = []
    for row in body:
        text = canonical_date(row[0])
        if text:
            current = text
        filled.append(current)

    dates: list[str] = []
    slot_texts: list[str] = []
    fc = np.empty((expected_days, len(hours), FC_HORIZON), dtype=float)
    for day in range(expected_days):
        day_dates = set()
        for slot_index, hour in enumerate(hours):
            offset = day * len(hours) + slot_index
            day_dates.add(filled[offset])
            got_hour = _slot_hour(body[offset][1])
            if got_hour != int(hour):
                raise InputValidationError(
                    f"附件 3 第 {offset + 2} 行预报时刻 {body[offset][1]!r} != 期望 {hour}:00"
                )
            try:
                values = np.array([float(item) for item in body[offset][2:]], dtype=float)
            except (TypeError, ValueError) as exc:
                raise InputValidationError(f"附件 3 第 {offset + 2} 行存在非数值预报：{exc}") from exc
            fc[day, slot_index, :] = values
        if len(day_dates) != 1 or not day_dates.pop():
            raise InputValidationError(f"附件 3 第 {day + 1} 天的日期前向填充不唯一")
        dates.append(filled[day * len(hours)])
        slot_texts.append(f"{int(hours[0])}:00")
    if not np.all(np.isfinite(fc)):
        raise InputValidationError("附件 3 预报含 NaN/Inf")
    if np.any(fc < 0):
        raise InputValidationError("附件 3 存在负预报")
    for index in range(1, len(dates)):
        if dates[index] == dates[index - 1]:
            raise InputValidationError(f"附件 3 日期未推进：{dates[index]}")
    return Attachment3(fc_kw=fc, dates=dates, slots=slot_texts, md5=md5_of(path))


def read_attachment4_price(
    path: Path, expected_days: int = DAYS_FULL, expected_periods: int = PERIODS_PER_DAY
) -> Attachment4:
    """读取附件 4 的 ``Sheet1``（366 行 × 145 列）→ ``(365, 144)`` 实际价矩阵。

    表头 ``日期\\时间`` + 144 个 10 分钟列，与附件 2 **同相位**（列数与标签逐列一致由调用方核对）。
    """
    if not path.is_file():
        raise InputValidationError(f"附件 4 不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        titles = list(workbook.sheetnames)
        if PRICE_SHEET not in titles:
            raise InputValidationError(f"附件 4 缺少工作表『{PRICE_SHEET}』：{titles}")
        sheet = workbook[PRICE_SHEET]
        rows = [row for row in sheet.iter_rows(min_row=1, max_col=expected_periods + 1, values_only=True)]
    finally:
        workbook.close()
    if len(rows) != expected_days + 1:
        raise InputValidationError(f"附件 4 行数 {len(rows)} != 表头 + {expected_days}")
    # 同 ``_read_power_sheet``：时刻列表头统一规范为 ``H:MM``，使附件 4 与附件 2 的
    # ``time_labels`` 逐列一致（相位核对）且与 ``AS01`` 左端点口径一致。
    header = ["" if item is None else _clock_label(item) for item in rows[0]]
    if header[0].strip() != POWER_DATE_HEADER:
        raise InputValidationError(f"附件 4 首列表头不是 {POWER_DATE_HEADER}：{header[0]!r}")
    if len(header) != expected_periods + 1:
        raise InputValidationError(f"附件 4 列数 {len(header)} != {expected_periods + 1}")
    body = rows[1:]
    dates = [_date_label(row[0]) for row in body]
    try:
        price = np.array([[float(item) for item in row[1:]] for row in body], dtype=float)
    except (TypeError, ValueError) as exc:
        raise InputValidationError(f"附件 4 存在非数值单元格：{exc}") from exc
    if price.shape != (expected_days, expected_periods):
        raise InputValidationError(f"附件 4 矩阵形状 {price.shape} 不符")
    if not np.all(np.isfinite(price)):
        raise InputValidationError("附件 4 含 NaN/Inf")
    if np.any(price < 0):
        raise InputValidationError("附件 4 存在负电价")
    _validate_dates(dates, source="附件 4 ")
    return Attachment4(
        price=price,
        dates=dates,
        time_labels=[str(item) for item in header[1:]],
        header=header,
        md5=md5_of(path),
    )


def _template_header(header: list[Any], sheet_name: str, *, periods: int = PERIODS_PER_DAY) -> list[str]:
    if len(header) != periods + 3:
        raise InputValidationError(f"模板『{sheet_name}』列数 {len(header)} != {periods + 3}")
    if str(header[0]).strip() != PLAN_HEADER_FIRST:
        raise InputValidationError(f"模板『{sheet_name}』首列表头异常：{header[0]!r}")
    if str(header[periods + 1]).strip() != PLAN_HEADER_ENERGY:
        raise InputValidationError(f"模板『{sheet_name}』第 {periods + 2} 列不是『{PLAN_HEADER_ENERGY}』")
    if str(header[periods + 2]).strip() != PLAN_HEADER_COST:
        raise InputValidationError(f"模板『{sheet_name}』第 {periods + 3} 列不是『{PLAN_HEADER_COST}』")
    return [str(item) for item in header[1 : periods + 1]]


def _inspect_template(
    path: Path,
    *,
    required_sheets: Sequence[str],
    purchase_sheets: Sequence[str],
    label_columns: int,
) -> dict[str, Any]:
    """只读打开模板并**实测**工作表集合、行列数、列标签、六块与 ``时刻`` 前两行、``⁝`` 压缩行。"""
    if not path.is_file():
        raise InputValidationError(f"模板不存在：{path}")
    measured: dict[str, Any] = {"sheets": [], "purchase": {}, "storage": {}, "emergency": {}}
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        titles = list(workbook.sheetnames)
        measured["sheets"] = titles
        for name in required_sheets:
            if name not in titles:
                raise InputValidationError(f"模板缺少工作表『{name}』：{titles}")
        for name in purchase_sheets:
            sheet = workbook[name]
            labels = _template_header([cell.value for cell in sheet[1]], name)
            dates = [_date_label(sheet.cell(row=row, column=1).value) for row in range(2, sheet.max_row + 1)]
            measured["purchase"][name] = {
                "labels": labels,
                "label_count": len(labels),
                "rows": len(dates),
                "first_date": dates[0] if dates else None,
                "last_date": dates[-1] if dates else None,
                "max_column": int(sheet.max_column),
            }
        storage = workbook[STORAGE_SHEET]
        storage_header = ["" if cell.value is None else str(cell.value) for cell in storage[1]]
        if tuple(storage_header[: len(STORAGE_HEADER)]) != STORAGE_HEADER:
            raise InputValidationError(f"模板『{STORAGE_SHEET}』表头异常：{storage_header}")
        block_rows = min(BLOCKS_PER_DAY, max(int(storage.max_row) - 1, 0))
        measured["storage"] = {
            "header": storage_header,
            "max_row": int(storage.max_row),
            "blocks": [str(storage.cell(row=2 + k, column=2).value) for k in range(block_rows)],
            "time_first_two_rows": [
                "" if storage.cell(row=2 + k, column=5).value is None else str(storage.cell(row=2 + k, column=5).value)
                for k in range(min(2, block_rows))
            ],
            "label_columns": label_columns,
        }
        emergency = workbook[EMERGENCY_SHEET]
        emergency_header = ["" if cell.value is None else str(cell.value) for cell in emergency[1]]
        if tuple(emergency_header[: len(EMERGENCY_HEADER)]) != EMERGENCY_HEADER:
            raise InputValidationError(f"模板『{EMERGENCY_SHEET}』表头异常：{emergency_header}")
        measured["emergency"] = {
            "header": emergency_header,
            "max_row": int(emergency.max_row),
            "has_compact_row": any(
                isinstance(cell, str) and "⁝" in cell
                for row in emergency.iter_rows(values_only=True)
                for cell in row
            ),
        }
        measured["merged"] = {name: [str(rng) for rng in workbook[name].merged_cells.ranges] for name in titles}
    finally:
        workbook.close()
    measured["md5"] = md5_of(path)
    return measured


def _geometry_differences(
    measured: dict[str, Any], expected: dict[str, Any], *, path: Path
) -> list[dict[str, Any]]:
    """把「实测 vs formulation §8.3 登记值」的差异逐条列出（实测以模板为准，差异只登记不修改）。"""
    diffs: list[dict[str, Any]] = []
    for key, want in expected.items():
        if key == "sheets":
            got = measured["sheets"]
        elif key.endswith(".rows"):
            got = measured["purchase"][key.split(".")[0]]["rows"]
        elif key.endswith(".labels_last"):
            got = measured["purchase"][key.split(".")[0]]["labels"][-1]
        elif key.endswith(".labels_first"):
            got = measured["purchase"][key.split(".")[0]]["labels"][0]
        elif key == "storage.max_row":
            got = measured["storage"]["max_row"]
        elif key == "emergency.max_row":
            got = measured["emergency"]["max_row"]
        else:
            continue
        if got != want:
            diffs.append({"field": key, "expected": want, "measured": got, "template": str(path)})
    return diffs


def inspect_template42(path: Path) -> dict[str, Any]:
    """只读复核 ``result4-2.xlsx``（计划购电量/充放电量/紧急购电量），实测值以模板为准。"""
    measured = _inspect_template(
        path,
        required_sheets=(PLAN_SHEET, STORAGE_SHEET, EMERGENCY_SHEET),
        purchase_sheets=(PLAN_SHEET,),
        label_columns=PERIODS_PER_DAY,
    )
    expected = {
        "sheets": [PLAN_SHEET, STORAGE_SHEET, EMERGENCY_SHEET],
        f"{PLAN_SHEET}.rows": 334,
        f"{PLAN_SHEET}.labels_first": "0:10-0:20",
        f"{PLAN_SHEET}.labels_last": "0:00-0:10+1",
        "storage.max_row": 20,
        "emergency.max_row": 11,
    }
    differences = _geometry_differences(measured, expected, path=path)
    return {
        "template": str(path),
        "measured": measured,
        "formulation_registered": expected,
        "differences": differences,
        "difference_note": (
            "formulation §8.3 的行数（充放电量 max_row / 紧急购电量 max_row）来自 prob02 模板示例，"
            "本问以**只读实测**为准并逐条登记；差异不代表模板损坏，也不修改模板。"
        ),
    }


def inspect_template43(path: Path) -> dict[str, Any]:
    """只读复核 ``result4-3.xlsx``（计划购电量/调整购电量/充放电量/紧急购电量），实测值以模板为准。"""
    measured = _inspect_template(
        path,
        required_sheets=(PLAN_SHEET, ADJUST_SHEET, STORAGE_SHEET, EMERGENCY_SHEET),
        purchase_sheets=(PLAN_SHEET, ADJUST_SHEET),
        label_columns=PERIODS_PER_DAY,
    )
    plan = measured["purchase"][PLAN_SHEET]
    adjust = measured["purchase"][ADJUST_SHEET]
    label_consistency = plan["labels"] == adjust["labels"]
    expected = {
        "sheets": [PLAN_SHEET, ADJUST_SHEET, STORAGE_SHEET, EMERGENCY_SHEET],
        f"{PLAN_SHEET}.rows": 334,
        f"{ADJUST_SHEET}.rows": 334,
        f"{PLAN_SHEET}.labels_first": "0:10-0:20",
        f"{PLAN_SHEET}.labels_last": "0:00-0:10+1",
        "storage.max_row": 26,
        "emergency.max_row": 11,
    }
    differences = _geometry_differences(measured, expected, path=path)
    if not label_consistency:
        differences.append(
            {"field": "plan_vs_adjust_labels", "expected": "identical", "measured": "different", "template": str(path)}
        )
    return {
        "template": str(path),
        "measured": measured,
        "formulation_registered": expected,
        "plan_adjust_labels_identical": bool(label_consistency),
        "differences": differences,
        "difference_note": (
            "formulation §8.3 的行数来自 prob03 模板示例，本问以**只读实测**为准并逐条登记；"
            "`计划购电量` 与 `调整购电量` 的列标签须逐列相同（替代量口径）。"
        ),
    }


def _assign(sheet: Any, *, row: int, column: int, value: Any) -> None:
    """显式写单元格，``value=None`` 也**真正清空**（防 openpyxl 静默 no-op）。"""
    sheet.cell(row=row, column=column).value = value


def _num(value: Any) -> float:
    """交付数值规整：保留 ``DELIVERY_DECIMALS`` 位小数。"""
    return float(round(float(value), DELIVERY_DECIMALS))


def _state(value: Any, *, clip_log: list[float]) -> float:
    """储电量交付值：先规整，再投影到 ``[1200, 10800]``，并把最大修正量记入 ``clip_log``。"""
    raw = float(value)
    rounded = round(raw, DELIVERY_DECIMALS)
    projected = min(max(rounded, STATE_MIN_KWH), STATE_MAX_KWH)
    clip_log.append(abs(projected - raw))
    return float(projected)


def _clear_rows(sheet: Any, *, first_row: int, last_row: int, last_column: int) -> None:
    for row in range(first_row, last_row + 1):
        for column in range(1, last_column + 1):
            _assign(sheet, row=row, column=column, value=None)


def _residual_report(
    *, workbook: Any, purchase_sheets: Sequence[str], days: int, emergency_rows: int
) -> list[str]:
    """复核填报后的工作簿：枚举**应空却仍有值**的单元格（防静默 no-op 复发）。"""
    residuals: list[str] = []
    for name in purchase_sheets:
        sheet = workbook[name]
        for row in range(days + 2, int(sheet.max_row) + 1):
            for column in range(1, PERIODS_PER_DAY + 4):
                if sheet.cell(row=row, column=column).value is not None:
                    residuals.append(f"{name}!{get_column_letter(column)}{row}（超出行数）")
                    break
    storage = workbook[STORAGE_SHEET]
    for index in range(days):
        for block in range(BLOCKS_PER_DAY):
            row = 2 + index * BLOCKS_PER_DAY + block
            if block != 0 and storage.cell(row=row, column=1).value is not None:
                residuals.append(f"充放电量!A{row}（日期只写块首行）")
            if block >= 2 and storage.cell(row=row, column=5).value is not None:
                residuals.append(f"充放电量!E{row}（时刻只写每日前两行）")
            if block >= 2 and storage.cell(row=row, column=6).value is not None:
                residuals.append(f"充放电量!F{row}（储电量只写每日前两行）")
    expected_storage_rows = 1 + days * BLOCKS_PER_DAY
    for row in range(expected_storage_rows + 1, int(storage.max_row) + 1):
        for column in range(1, len(STORAGE_HEADER) + 1):
            if storage.cell(row=row, column=column).value is not None:
                residuals.append(f"充放电量!{get_column_letter(column)}{row}（超出行数）")
                break
    emergency = workbook[EMERGENCY_SHEET]
    for row in range(1 + emergency_rows + 1, int(emergency.max_row) + 1):
        for column in range(1, len(EMERGENCY_HEADER) + 1):
            if emergency.cell(row=row, column=column).value is not None:
                residuals.append(f"紧急购电量!{get_column_letter(column)}{row}（超出行数）")
                break
    for name in workbook.sheetnames:
        for row in workbook[name].iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str) and "⁝" in value:
                    residuals.append(f"{name}!不得残留模板 ⁝ 压缩行")
    total = len(residuals)
    if total > 200:
        residuals = residuals[:200] + [f"…（共 {total} 处，已截断）"]
    return residuals


def _fill_purchase_sheet(
    sheet: Any,
    *,
    name: str,
    dates: Sequence[str],
    values: np.ndarray,
    day_energy_kwh: np.ndarray,
    day_cost_yuan: np.ndarray,
) -> None:
    days = len(dates)
    if sheet.max_row < days + 1:
        raise InputValidationError(f"模板『{name}』只有 {sheet.max_row - 1} 行，少于 {days} 天")
    for index, date_text in enumerate(dates):
        row = 2 + index
        actual = _date_label(sheet.cell(row=row, column=1).value)
        if actual != date_text:
            raise InputValidationError(f"模板『{name}』第 {row} 行日期 {actual!r} != {date_text!r}")
        for period in range(PERIODS_PER_DAY):
            _assign(sheet, row=row, column=2 + period, value=_num(values[index, period]))
        _assign(sheet, row=row, column=PERIODS_PER_DAY + 2, value=_num(day_energy_kwh[index]))
        _assign(sheet, row=row, column=PERIODS_PER_DAY + 3, value=_num(day_cost_yuan[index]))
    if sheet.max_row > days + 1:
        _clear_rows(sheet, first_row=days + 2, last_row=sheet.max_row, last_column=PERIODS_PER_DAY + 3)


def _fill_storage_sheet(
    storage: Any,
    *,
    dates: Sequence[str],
    charge: np.ndarray,
    discharge: np.ndarray,
    state_start_kwh: np.ndarray,
    state_end_kwh: np.ndarray,
    clip_log: list[float],
) -> tuple[int, list[str]]:
    blocks = [str(storage.cell(row=2 + k, column=2).value) for k in range(BLOCKS_PER_DAY)]
    filled = 0
    for index, date_text in enumerate(dates):
        date_value = datetime.strptime(date_text, "%Y-%m-%d")
        for block in range(BLOCKS_PER_DAY):
            row = 2 + index * BLOCKS_PER_DAY + block
            start = block * BLOCK_ROWS
            stop = start + BLOCK_ROWS
            _assign(storage, row=row, column=1, value=date_value if block == 0 else None)
            _assign(storage, row=row, column=2, value=blocks[block])
            _assign(storage, row=row, column=3, value=_num(np.sum(charge[index, start:stop])))
            _assign(storage, row=row, column=4, value=_num(np.sum(discharge[index, start:stop])))
            _assign(
                storage,
                row=row,
                column=5,
                value="00:00:00" if block == 0 else ("24:00" if block == 1 else None),
            )
            _assign(
                storage,
                row=row,
                column=6,
                value=_state(state_start_kwh[index], clip_log=clip_log) if block == 0 else (
                    _state(state_end_kwh[index], clip_log=clip_log) if block == 1 else None
                ),
            )
            filled += 1
    expected_rows = 1 + len(dates) * BLOCKS_PER_DAY
    if storage.max_row > expected_rows:
        _clear_rows(storage, first_row=expected_rows + 1, last_row=storage.max_row, last_column=6)
    return filled, blocks


def _fill_emergency_sheet(
    emergency: Any,
    *,
    dates: Sequence[str],
    emergency_intervals: Sequence[Sequence[tuple[str, float]]],
) -> int:
    row = 2
    filled = 0
    for index, date_text in enumerate(dates):
        date_value = datetime.strptime(date_text, "%Y-%m-%d")
        intervals = list(emergency_intervals[index])
        if not intervals:
            intervals = [(ZERO_INTERVAL_TEXT, 0.0)]
        for order, (slot, kwh) in enumerate(intervals):
            _assign(emergency, row=row, column=1, value=date_value if order == 0 else None)
            _assign(emergency, row=row, column=2, value=slot)
            _assign(emergency, row=row, column=3, value=_num(kwh))
            row += 1
            filled += 1
    expected_rows = row - 1
    if emergency.max_row > expected_rows:
        _clear_rows(emergency, first_row=expected_rows + 1, last_row=emergency.max_row, last_column=3)
    return filled


def fill_result42_workbook(
    *,
    template: Path,
    dest: Path,
    dates: Sequence[str],
    purchase_kwh: np.ndarray,
    charge_kwh: np.ndarray,
    discharge_kwh: np.ndarray,
    state_start_kwh: np.ndarray,
    state_end_kwh: np.ndarray,
    day_energy_kwh: np.ndarray,
    day_cost_yuan: np.ndarray,
    emergency_intervals: Sequence[Sequence[tuple[str, float]]],
) -> dict[str, Any]:
    """把 ``4-2`` 链的交付期解填入 ``result4-2.xlsx``（三工作表）。

    口径（AS20 + formulation §8.3）：
      * ``计划购电量``：``b``；末两列 ``全天购电量 = Σb``、``全天购电费 = C_total^{act}``（``4-2`` 无 ``C_adj``）；
      * ``充放电量``：334 天 × 6 个 4 小时块，``充电量``/``放电量`` **分别列示不冲抵**，日期只写块首行，
        ``时刻`` = ``00:00:00``/``24:00`` 与 ``储电量`` = ``E_{d,0}``/``E_{d,144}`` 只写每日前两行；
      * ``紧急购电量``：每日至少 1 行；``J_d = 0`` 写 ``—``/``0``。
    """
    days = len(dates)
    purchase = np.asarray(purchase_kwh, dtype=float)
    charge = np.asarray(charge_kwh, dtype=float)
    discharge = np.asarray(discharge_kwh, dtype=float)
    clip_log: list[float] = []
    if purchase.shape != (days, PERIODS_PER_DAY):
        raise InputValidationError(f"purchase 形状 {purchase.shape} != ({days}, {PERIODS_PER_DAY})")
    for label, array in (("charge", charge), ("discharge", discharge)):
        if array.shape != purchase.shape:
            raise InputValidationError(f"{label} 形状与 purchase 不一致")
    if len(emergency_intervals) != days:
        raise InputValidationError("emergency_intervals 天数与 dates 不一致")
    if not template.is_file():
        raise InputValidationError(f"模板不存在：{template}")

    workbook = openpyxl.load_workbook(template)
    try:
        titles = list(workbook.sheetnames)
        for name in (PLAN_SHEET, STORAGE_SHEET, EMERGENCY_SHEET):
            if name not in titles:
                raise InputValidationError(f"模板缺少工作表『{name}』：{titles}")
        _fill_purchase_sheet(
            workbook[PLAN_SHEET],
            name=PLAN_SHEET,
            dates=dates,
            values=purchase,
            day_energy_kwh=day_energy_kwh,
            day_cost_yuan=day_cost_yuan,
        )
        blocks_filled, blocks = _fill_storage_sheet(
            workbook[STORAGE_SHEET],
            dates=dates,
            charge=charge,
            discharge=discharge,
            state_start_kwh=state_start_kwh,
            state_end_kwh=state_end_kwh,
            clip_log=clip_log,
        )
        emergency_rows = _fill_emergency_sheet(
            workbook[EMERGENCY_SHEET], dates=dates, emergency_intervals=emergency_intervals
        )
        dest.parent.mkdir(parents=True, exist_ok=True)
        residuals = _residual_report(
            workbook=workbook, purchase_sheets=(PLAN_SHEET,), days=days, emergency_rows=emergency_rows
        )
        workbook.save(dest)
    finally:
        workbook.close()
    return {
        "chain": "4-2",
        "plan_rows_filled": days,
        "plan_periods_per_row": PERIODS_PER_DAY,
        "storage_blocks_filled": blocks_filled,
        "storage_blocks": blocks,
        "emergency_rows_filled": emergency_rows,
        "format_residuals": residuals,
        "delivery_decimals": DELIVERY_DECIMALS,
        "storage_projection": {
            "bounds_kwh": [STATE_MIN_KWH, STATE_MAX_KWH],
            "max_abs_correction_kwh": max(clip_log) if clip_log else 0.0,
            "cells": len(clip_log),
        },
    }


def fill_result43_workbook(
    *,
    template: Path,
    dest: Path,
    dates: Sequence[str],
    planned_kwh: np.ndarray,
    adjusted_kwh: np.ndarray,
    charge_kwh: np.ndarray,
    discharge_kwh: np.ndarray,
    state_start_kwh: np.ndarray,
    state_end_kwh: np.ndarray,
    day_energy_planned_kwh: np.ndarray,
    day_energy_adjusted_kwh: np.ndarray,
    day_cost_yuan: np.ndarray,
    emergency_intervals: Sequence[Sequence[tuple[str, float]]],
) -> dict[str, Any]:
    """把 ``4-3`` 链的交付期解填入 ``result4-3.xlsx``（四工作表）。

    口径（AS20 + formulation §8.3）：
      * ``计划购电量``：``b``；末两列 ``Σb``、``C_total^{act}``；
      * ``调整购电量``：**最终量** ``q``（替代量口径）；末两列 ``Σq``、``C_total^{act}``；
      * ``充放电量`` / ``紧急购电量``：与 ``4-2`` 同款几何。
    """
    days = len(dates)
    planned = np.asarray(planned_kwh, dtype=float)
    adjusted = np.asarray(adjusted_kwh, dtype=float)
    charge = np.asarray(charge_kwh, dtype=float)
    discharge = np.asarray(discharge_kwh, dtype=float)
    clip_log: list[float] = []
    if planned.shape != (days, PERIODS_PER_DAY):
        raise InputValidationError(f"planned 形状 {planned.shape} != ({days}, {PERIODS_PER_DAY})")
    for label, array in (("adjusted", adjusted), ("charge", charge), ("discharge", discharge)):
        if array.shape != planned.shape:
            raise InputValidationError(f"{label} 形状与 planned 不一致")
    if len(emergency_intervals) != days:
        raise InputValidationError("emergency_intervals 天数与 dates 不一致")
    if not template.is_file():
        raise InputValidationError(f"模板不存在：{template}")

    workbook = openpyxl.load_workbook(template)
    try:
        titles = list(workbook.sheetnames)
        for name in (PLAN_SHEET, ADJUST_SHEET, STORAGE_SHEET, EMERGENCY_SHEET):
            if name not in titles:
                raise InputValidationError(f"模板缺少工作表『{name}』：{titles}")
        for name, values, energy in (
            (PLAN_SHEET, planned, day_energy_planned_kwh),
            (ADJUST_SHEET, adjusted, day_energy_adjusted_kwh),
        ):
            _fill_purchase_sheet(
                workbook[name],
                name=name,
                dates=dates,
                values=values,
                day_energy_kwh=energy,
                day_cost_yuan=day_cost_yuan,
            )
        blocks_filled, blocks = _fill_storage_sheet(
            workbook[STORAGE_SHEET],
            dates=dates,
            charge=charge,
            discharge=discharge,
            state_start_kwh=state_start_kwh,
            state_end_kwh=state_end_kwh,
            clip_log=clip_log,
        )
        emergency_rows = _fill_emergency_sheet(
            workbook[EMERGENCY_SHEET], dates=dates, emergency_intervals=emergency_intervals
        )
        dest.parent.mkdir(parents=True, exist_ok=True)
        residuals = _residual_report(
            workbook=workbook,
            purchase_sheets=(PLAN_SHEET, ADJUST_SHEET),
            days=days,
            emergency_rows=emergency_rows,
        )
        workbook.save(dest)
    finally:
        workbook.close()
    return {
        "chain": "4-3",
        "plan_rows_filled": days,
        "adjust_rows_filled": days,
        "plan_periods_per_row": PERIODS_PER_DAY,
        "storage_blocks_filled": blocks_filled,
        "storage_blocks": blocks,
        "emergency_rows_filled": emergency_rows,
        "format_residuals": residuals,
        "delivery_decimals": DELIVERY_DECIMALS,
        "storage_projection": {
            "bounds_kwh": [STATE_MIN_KWH, STATE_MAX_KWH],
            "max_abs_correction_kwh": max(clip_log) if clip_log else 0.0,
            "cells": len(clip_log),
        },
    }


def write_json(path: Path, payload: Any) -> None:
    """写 JSON；显式拒绝 NaN/Inf，避免把非有限值伪装成合法结果。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False)
    path.write_text(text + "\n", encoding="utf-8")


def delivery_dates(all_dates: Sequence[str], start_day_index: int) -> list[str]:
    """交付期日期列表（默认 2025-02-01 … 2025-12-31，334 天）。"""
    return [str(item) for item in all_dates[start_day_index:]]


def day_offset(dates: Sequence[str], target: str) -> int:
    """在日期序列中定位目标的 0 基日索引；找不到时严格失败。"""
    for index, item in enumerate(dates):
        if item == target:
            return index
    raise InputValidationError(f"日期序列中找不到 {target}")


def _format_minutes(total_minutes: int) -> str:
    """把「左端点分钟数」格式化为 ``H:MM``；≥1440 时按 AS01 写 ``+1``。"""
    if total_minutes < 1440:
        return f"{total_minutes // 60}:{total_minutes % 60:02d}"
    shifted = total_minutes - 1440
    return f"{shifted // 60}:{shifted % 60:02d}+1"


def interval_label(start: int, stop: int) -> str:
    """按 AS01 生成 ``[τ_start+1, τ_stop+1 + 10min)`` 的左端点区间标签。

    ``start``/``stop`` 为 0 基时段下标（半开区间 ``[start, stop)``）；两端**都**用 ``H:MM`` 风格
    （小时不补零），与附件 5 模板列标签风格（如 ``9:50-10:00``）一致（承 prob03 勘误 R9）。
    """
    return f"{_format_minutes(10 * (start + 1))}-{_format_minutes(10 * (stop + 1))}"


__all__ = [
    "ADJUST_SHEET",
    "Attachment2",
    "Attachment3",
    "Attachment4",
    "BLOCKS_PER_DAY",
    "BLOCK_ROWS",
    "DAYS_FULL",
    "DECISION_HOURS",
    "DELIVERY_DECIMALS",
    "EMERGENCY_SHEET",
    "FC_HORIZON",
    "InputValidationError",
    "LOAD_SHEET",
    "PERIODS_PER_DAY",
    "PLAN_SHEET",
    "PRICE_SHEET",
    "PV_SHEET",
    "STORAGE_SHEET",
    "ZERO_INTERVAL_TEXT",
    "canonical_date",
    "day_offset",
    "delivery_dates",
    "fill_result42_workbook",
    "fill_result43_workbook",
    "inspect_template42",
    "inspect_template43",
    "interval_label",
    "md5_of",
    "read_attachment2",
    "read_attachment3",
    "read_attachment4_price",
    "write_json",
]
