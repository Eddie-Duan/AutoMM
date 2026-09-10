# -*- coding: utf-8 -*-
"""prob01 输入读取与 ``result1.xlsx`` 填报。

只读 ``data/附件1.xlsx`` 与模板 ``data/附件5/result1.xlsx``；不修改任何原始数据。
口径来源：assumption_v003（AS01/AS02/AS09/AS13）与团队勘误 E1/E2/E3。
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import openpyxl

EXPECTED_HEADER = ("时间", "电价", "小区负载", "光伏发电预测功率")
PLAN_SHEET = "计划购电量"
STORAGE_SHEET = "充放电量"
BLOCK_ROWS = 24


class InputValidationError(RuntimeError):
    """输入结构与 accepted 假设不一致时严格失败。"""


@dataclass(frozen=True)
class Attachment1:
    """附件 1 的 144 时段电价 / 负载 / 光伏预测（功率单位 kW）。"""

    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray
    time_labels: list[str]
    md5: str
    header: list[str]

    @property
    def periods(self) -> int:
        return int(self.price.shape[0])


def md5_of(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _label(value: Any) -> str:
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    return str(value)


def read_attachment1(path: Path, expected_rows: int = 144) -> Attachment1:
    """读取附件 1 第 1 个工作表并做结构校验。"""
    if not path.is_file():
        raise InputValidationError(f"附件 1 不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [row for row in sheet.iter_rows(min_row=1, max_col=4, values_only=True)]
    finally:
        workbook.close()
    if len(rows) < 2:
        raise InputValidationError("附件 1 没有数据行")
    header = [str(item).strip() for item in rows[0]]
    if tuple(header) != EXPECTED_HEADER:
        raise InputValidationError(f"附件 1 表头不匹配：{header} != {list(EXPECTED_HEADER)}")
    body = rows[1:]
    if len(body) != expected_rows:
        raise InputValidationError(f"附件 1 数据行数 {len(body)} != 期望 {expected_rows}")
    labels = [_label(row[0]) for row in body]
    try:
        price = np.array([float(row[1]) for row in body], dtype=float)
        load_kw = np.array([float(row[2]) for row in body], dtype=float)
        pv_kw = np.array([float(row[3]) for row in body], dtype=float)
    except (TypeError, ValueError) as exc:
        raise InputValidationError(f"附件 1 存在非数值单元格：{exc}") from exc
    for name, values in (("电价", price), ("负载", load_kw), ("光伏", pv_kw)):
        if not np.all(np.isfinite(values)):
            raise InputValidationError(f"附件 1 的{name}列含 NaN/Inf")
    if np.any(price <= 0):
        raise InputValidationError("附件 1 存在非正电价，LP 对偶与互补性引理的前提被破坏")
    if np.any(load_kw < 0):
        raise InputValidationError("附件 1 存在负负载")
    if np.any(pv_kw < 0):
        raise InputValidationError("附件 1 存在负光伏功率")
    return Attachment1(
        price=price,
        load_kw=load_kw,
        pv_kw=pv_kw,
        time_labels=labels,
        md5=md5_of(path),
        header=header,
    )


def inspect_template(path: Path) -> dict[str, Any]:
    """读取模板的逐行标签，用于 AS01/AS13 的对齐校验（不写入）。"""
    if not path.is_file():
        raise InputValidationError(f"模板不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if PLAN_SHEET not in workbook.sheetnames or STORAGE_SHEET not in workbook.sheetnames:
            raise InputValidationError(f"模板工作表名不符合要求：{workbook.sheetnames}")
        plan = workbook[PLAN_SHEET]
        plan_labels = [str(row[0]) for row in plan.iter_rows(min_row=2, max_col=1, values_only=True)]
        storage = workbook[STORAGE_SHEET]
        storage_rows = [row for row in storage.iter_rows(min_row=2, max_col=5, values_only=True)]
    finally:
        workbook.close()
    return {
        "plan_labels": plan_labels,
        "storage_labels": [str(row[0]) for row in storage_rows],
        "storage_endpoint_labels": [str(row[3]) if row[3] is not None else None for row in storage_rows[:2]],
        "plan_periods": len(plan_labels),
    }


def fill_result_workbook(
    *,
    template: Path,
    dest: Path,
    purchase_kwh: Sequence[float],
    charge_kwh: Sequence[float],
    discharge_kwh: Sequence[float],
    e_init: float,
    e_final: float,
    periods: int,
) -> dict[str, Any]:
    """把解填入 ``result1.xlsx``；保留模板标签与样式，只写数值单元格。"""
    if not template.is_file():
        raise InputValidationError(f"模板不存在：{template}")
    workbook = openpyxl.load_workbook(template)
    try:
        if PLAN_SHEET not in workbook.sheetnames or STORAGE_SHEET not in workbook.sheetnames:
            raise InputValidationError(f"模板工作表名不符合要求：{workbook.sheetnames}")
        plan = workbook[PLAN_SHEET]
        plan_rows = plan.max_row - 1
        if plan_rows < periods:
            raise InputValidationError(f"模板『{PLAN_SHEET}』只有 {plan_rows} 行，少于 {periods} 个时段")
        for index in range(periods):
            plan.cell(row=2 + index, column=2, value=float(purchase_kwh[index]))

        storage = workbook[STORAGE_SHEET]
        blocks = storage.max_row - 1
        filled_blocks = 0
        for block in range(blocks):
            start = block * BLOCK_ROWS
            stop = min(start + BLOCK_ROWS, periods)
            if stop <= start:
                break
            storage.cell(row=2 + block, column=2, value=float(sum(charge_kwh[start:stop])))
            storage.cell(row=2 + block, column=3, value=float(sum(discharge_kwh[start:stop])))
            if stop - start == BLOCK_ROWS:
                filled_blocks += 1
        if periods == blocks * BLOCK_ROWS:
            storage.cell(row=2, column=5, value=float(e_init))
            storage.cell(row=3, column=5, value=float(e_final))
        dest.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(dest)
    finally:
        workbook.close()
    return {
        "plan_rows_filled": periods,
        "storage_blocks_filled": filled_blocks,
        "storage_endpoints_filled": periods == 144,
    }


def write_json(path: Path, payload: Any) -> None:
    """写 JSON；显式拒绝 NaN/Inf，避免把非有限值伪装成合法结果。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False)
    path.write_text(text + "\n", encoding="utf-8")
