# -*- coding: utf-8 -*-
"""prob02 输入读取与 ``result2.xlsx`` 填报。

只读 ``data/附件1.xlsx``、``data/附件2.xlsx`` 与模板 ``data/附件5/result2.xlsx``；不修改任何原始数据。
口径来源：prob02 assumption_v001（AS01/AS02/AS12 + 团队勘误 R1/R3/R4）与 formulation_v001 §7。

纪律（AS02 / B4）：附件 1 **只使用电价列**；其负载/光伏预测列、附件 3、附件 4 均不进入模型。
本模块读取附件 1 时只为「表头结构校验」读取第 3/4 列表头文本，**不把这两列的数值传入模型**。
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

EXPECTED_HEADER_A1 = ("时间", "电价", "小区负载", "光伏发电预测功率")
LOAD_SHEET = "小区负载"
PV_SHEET = "光伏发电实际功率"
PLAN_SHEET = "计划购电量"
STORAGE_SHEET = "充放电量"
EMERGENCY_SHEET = "紧急购电量"
PLAN_HEADER_FIRST = "日期\\时间"
PLAN_HEADER_ENERGY = "全天购电量"
PLAN_HEADER_COST = "全天购电费"
STORAGE_HEADER = ("日期", "时间段", "充电量", "放电量", "时刻", "储电量")
EMERGENCY_HEADER = ("日期", "购电时间段", "购电量")
BLOCK_ROWS = 24
BLOCKS_PER_DAY = 6
PERIODS_PER_DAY = 144
ZERO_INTERVAL_TEXT = "—"
A2_DATE_HEADER = "日期\\时间"


class InputValidationError(RuntimeError):
    """输入结构与 accepted 假设不一致时严格失败。"""


@dataclass(frozen=True)
class Attachment1:
    """附件 1 的 144 个时段电价（元/kWh）与时间标签（只含电价列）。"""

    price: np.ndarray
    time_labels: list[str]
    header: list[str]
    md5: str

    @property
    def periods(self) -> int:
        return int(self.price.shape[0])


@dataclass(frozen=True)
class Attachment2:
    """附件 2 的逐日实际负载/光伏功率（365 天 × 144 个 10 分钟点，单位 kW）。"""

    load_kw: np.ndarray
    pv_kw: np.ndarray
    dates: list[str]
    time_labels: list[str]
    load_md5: str
    pv_md5: str
    md5: str

    @property
    def days(self) -> int:
        return int(self.load_kw.shape[0])

    @property
    def periods_per_day(self) -> int:
        return int(self.load_kw.shape[1])


def md5_of(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clock_label(value: Any) -> str:
    """把附件时间戳格式化为 ``H:MM``（datetime/time 走 strftime，字符串原样返回）。"""
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    return str(value).strip()


def _date_label(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def read_attachment1(path: Path, expected_rows: int = PERIODS_PER_DAY) -> Attachment1:
    """读取附件 1 第 1 个工作表；只保留电价列与时间标签。"""
    if not path.is_file():
        raise InputValidationError(f"附件 1 不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [row for row in sheet.iter_rows(min_row=1, max_col=4, values_only=True)]
    finally:
        workbook.close()
    if len(rows) < 2:
        raise InputValidationError("附件 1 没有数据行")
    header = [str(item).strip() for item in rows[0]]
    if tuple(header) != EXPECTED_HEADER_A1:
        raise InputValidationError(f"附件 1 表头不匹配：{header} != {list(EXPECTED_HEADER_A1)}")
    body = rows[1:]
    if len(body) != expected_rows:
        raise InputValidationError(f"附件 1 数据行数 {len(body)} != 期望 {expected_rows}")
    labels = [_clock_label(row[0]) for row in body]
    try:
        price = np.array([float(row[1]) for row in body], dtype=float)
    except (TypeError, ValueError) as exc:
        raise InputValidationError(f"附件 1 电价列存在非数值单元格：{exc}") from exc
    if not np.all(np.isfinite(price)):
        raise InputValidationError("附件 1 电价列含 NaN/Inf")
    if np.any(price <= 0):
        raise InputValidationError("附件 1 存在非正电价：定理 T1 与 LP 对偶前提被破坏")
    return Attachment1(price=price, time_labels=labels, header=header, md5=md5_of(path))


def _read_power_sheet(
    sheet: Any, *, name: str, expected_days: int, expected_periods: int
) -> tuple[list[str], list[str], np.ndarray]:
    rows = [row for row in sheet.iter_rows(min_row=1, max_col=expected_periods + 1, values_only=True)]
    if len(rows) != expected_days + 1:
        raise InputValidationError(f"附件 2『{name}』行数 {len(rows)} != 表头 + {expected_days}")
    header = ["" if item is None else str(item) for item in rows[0]]
    if header[0].strip() != A2_DATE_HEADER:
        raise InputValidationError(f"附件 2『{name}』首列表头不是 {A2_DATE_HEADER}：{header[0]!r}")
    if len(header) != expected_periods + 1:
        raise InputValidationError(f"附件 2『{name}』列数 {len(header)} != {expected_periods + 1}")
    body = rows[1:]
    dates = [_date_label(row[0]) for row in body]
    try:
        values = np.array([[float(item) for item in row[1:]] for row in body], dtype=float)
    except (TypeError, ValueError) as exc:
        raise InputValidationError(f"附件 2『{name}』存在非数值单元格：{exc}") from exc
    if values.shape != (expected_days, expected_periods):
        raise InputValidationError(f"附件 2『{name}』矩阵形状 {values.shape} 不符")
    if not np.all(np.isfinite(values)):
        raise InputValidationError(f"附件 2『{name}』含 NaN/Inf")
    if np.any(values < 0):
        raise InputValidationError(f"附件 2『{name}』存在负功率")
    return header, dates, values


def read_attachment2(
    path: Path, expected_days: int = 365, expected_periods: int = PERIODS_PER_DAY
) -> Attachment2:
    """读取附件 2 的『小区负载』与『光伏发电实际功率』（均为实际值）。"""
    if not path.is_file():
        raise InputValidationError(f"附件 2 不存在：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        titles = list(workbook.sheetnames)
        for name in (LOAD_SHEET, PV_SHEET):
            if name not in titles:
                raise InputValidationError(f"附件 2 缺少工作表『{name}』：{titles}")
        load_header, load_dates, load = _read_power_sheet(
            workbook[LOAD_SHEET], name=LOAD_SHEET, expected_days=expected_days, expected_periods=expected_periods
        )
        pv_header, pv_dates, pv = _read_power_sheet(
            workbook[PV_SHEET], name=PV_SHEET, expected_days=expected_days, expected_periods=expected_periods
        )
    finally:
        workbook.close()
    if load_dates != pv_dates:
        raise InputValidationError("附件 2 两个工作表的日期序列不一致")
    _validate_dates(load_dates)
    if load_header[1:] != pv_header[1:]:
        raise InputValidationError("附件 2 两个工作表的时间标签不一致")
    return Attachment2(
        load_kw=load,
        pv_kw=pv,
        dates=load_dates,
        time_labels=[str(item) for item in load_header[1:]],
        load_md5=md5_of(path),
        pv_md5=md5_of(path),
        md5=md5_of(path),
    )


def _validate_dates(dates: Sequence[str]) -> None:
    try:
        parsed = [datetime.strptime(item, "%Y-%m-%d").date() for item in dates]
    except ValueError as exc:
        raise InputValidationError(f"附件 2 日期列无法解析：{exc}") from exc
    for index in range(1, len(parsed)):
        if (parsed[index] - parsed[index - 1]).days != 1:
            raise InputValidationError(f"附件 2 日期不连续：{dates[index - 1]} -> {dates[index]}")
    if parsed[0].isoformat() != "2025-01-01":
        raise InputValidationError(f"附件 2 首日不是 2025-01-01：{dates[0]}")


def inspect_template(path: Path) -> dict[str, Any]:
    """读取 ``result2.xlsx`` 模板的逐行标签/表头，用于 AS01/AS12 对齐校验（不写入）。"""
    if not path.is_file():
        raise InputValidationError(f"模板不存在：{path}")
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        titles = list(workbook.sheetnames)
        for name in (PLAN_SHEET, STORAGE_SHEET, EMERGENCY_SHEET):
            if name not in titles:
                raise InputValidationError(f"模板缺少工作表『{name}』：{titles}")
        plan = workbook[PLAN_SHEET]
        plan_header = [cell.value for cell in plan[1]]
        if len(plan_header) < PERIODS_PER_DAY + 3:
            raise InputValidationError(f"模板『{PLAN_SHEET}』列数 {len(plan_header)} 少于 147")
        if str(plan_header[0]).strip() != PLAN_HEADER_FIRST:
            raise InputValidationError(f"模板『{PLAN_SHEET}』首列表头异常：{plan_header[0]!r}")
        if str(plan_header[PERIODS_PER_DAY + 1]).strip() != PLAN_HEADER_ENERGY:
            raise InputValidationError(f"模板『{PLAN_SHEET}』第 146 列不是『{PLAN_HEADER_ENERGY}』")
        if str(plan_header[PERIODS_PER_DAY + 2]).strip() != PLAN_HEADER_COST:
            raise InputValidationError(f"模板『{PLAN_SHEET}』第 147 列不是『{PLAN_HEADER_COST}』")
        plan_labels = [str(item) for item in plan_header[1 : PERIODS_PER_DAY + 1]]
        plan_dates = [_date_label(plan.cell(row=row, column=1).value) for row in range(2, plan.max_row + 1)]

        storage = workbook[STORAGE_SHEET]
        storage_header = ["" if cell.value is None else str(cell.value) for cell in storage[1]]
        if tuple(storage_header[: len(STORAGE_HEADER)]) != STORAGE_HEADER:
            raise InputValidationError(f"模板『{STORAGE_SHEET}』表头异常：{storage_header}")
        storage_blocks = [str(storage.cell(row=2 + k, column=2).value) for k in range(BLOCKS_PER_DAY)]
        storage_max_row = int(storage.max_row)

        emergency = workbook[EMERGENCY_SHEET]
        emergency_header = ["" if cell.value is None else str(cell.value) for cell in emergency[1]]
        if tuple(emergency_header[: len(EMERGENCY_HEADER)]) != EMERGENCY_HEADER:
            raise InputValidationError(f"模板『{EMERGENCY_SHEET}』表头异常：{emergency_header}")
        emergency_max_row = emergency.max_row
    finally:
        workbook.close()
    return {
        "plan_header_len": len(plan_header),
        "plan_labels": plan_labels,
        "plan_dates": plan_dates,
        "plan_rows": len(plan_dates),
        "storage_header": storage_header,
        "storage_blocks": storage_blocks,
        "storage_max_row": storage_max_row,
        "emergency_header": emergency_header,
        "emergency_max_row": emergency_max_row,
        "md5": md5_of(path),
    }


def _storage_max_row(path: Path) -> int:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return int(workbook[STORAGE_SHEET].max_row)
    finally:
        workbook.close()


def _assign(sheet: Any, *, row: int, column: int, value: Any) -> None:
    """显式写单元格，``value=None`` 也**真正清空**。

    缺陷登记（sanity act-ca831716e72246eb）：``Worksheet.cell(row, column, value=None)`` 在
    ``value is None`` 时**不赋值**，是静默 no-op，无法清空模板残留值（曾导致交付的
    ``result2.xlsx`` 残留 ``充放电量!A15=2025-12-31``、``!E16='24:00'``）。本函数统一改为
    直接设置 ``cell.value``，禁止再用 ``cell(..., value=None)`` 表达清空。
    """
    sheet.cell(row=row, column=column).value = value


def _clear_rows(sheet: Any, *, first_row: int, last_row: int, last_column: int) -> None:
    for row in range(first_row, last_row + 1):
        for column in range(1, last_column + 1):
            _assign(sheet, row=row, column=column, value=None)


def _residual_report(*, workbook: Any, days: int, emergency_rows: int) -> list[str]:
    """复核填报后的工作簿：枚举**应空却仍有值**的单元格（防静默 no-op 复发）。

    口径同 AS12 + 团队勘误 R3/R4：``充放电量`` 日期只写块首行、``时刻``/``储电量`` 只写每日前两行；
    ``计划购电量``/``紧急购电量`` 超出应有行数后不得留值；全簿不得残留模板 ``⁝`` 压缩行。
    """
    residuals: list[str] = []
    plan = workbook[PLAN_SHEET]
    for row in range(days + 2, int(plan.max_row) + 1):
        for column in range(1, PERIODS_PER_DAY + 4):
            if plan.cell(row=row, column=column).value is not None:
                residuals.append(f"计划购电量!{get_column_letter(column)}{row}")
                break
    storage = workbook[STORAGE_SHEET]
    for index in range(days):
        for block in range(BLOCKS_PER_DAY):
            row = 2 + index * BLOCKS_PER_DAY + block
            if block != 0 and storage.cell(row=row, column=1).value is not None:
                residuals.append(f"充放电量!A{row}（日期只写块首行）")
            if block >= 2 and storage.cell(row=row, column=5).value is not None:
                residuals.append(f"充放电量!E{row}（时刻只写每日前两行）")
            if block >= 2 and storage.cell(row=row, column=6).value is not None:
                residuals.append(f"充放电量!F{row}（储电量只写每日前两行）")
    expected_storage_rows = 1 + days * BLOCKS_PER_DAY
    for row in range(expected_storage_rows + 1, int(storage.max_row) + 1):
        for column in range(1, len(STORAGE_HEADER) + 1):
            if storage.cell(row=row, column=column).value is not None:
                residuals.append(f"充放电量!{get_column_letter(column)}{row}（超出行数）")
                break
    emergency = workbook[EMERGENCY_SHEET]
    for row in range(1 + emergency_rows + 1, int(emergency.max_row) + 1):
        for column in range(1, len(EMERGENCY_HEADER) + 1):
            if emergency.cell(row=row, column=column).value is not None:
                residuals.append(f"紧急购电量!{get_column_letter(column)}{row}（超出行数）")
                break
    for name in workbook.sheetnames:
        for row in workbook[name].iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str) and "⁝" in value:
                    residuals.append(f"{name}!不得残留模板 ⁝ 压缩行")
    total = len(residuals)
    if total > 200:
        residuals = residuals[:200] + [f"…（共 {total} 处，已截断）"]
    return residuals


def fill_result2_workbook(
    *,
    template: Path,
    dest: Path,
    dates: Sequence[str],
    purchase_kwh: np.ndarray,
    charge_kwh: np.ndarray,
    discharge_kwh: np.ndarray,
    state_start_kwh: np.ndarray,
    state_end_kwh: np.ndarray,
    day_energy_kwh: np.ndarray,
    day_cost_yuan: np.ndarray,
    emergency_intervals: Sequence[Sequence[tuple[str, float]]],
) -> dict[str, Any]:
    """把交付期解填入 ``result2.xlsx`` 三个工作表（保留模板标签与样式，只写数值单元格）。

    口径（AS12 + 勘误 R3/R4）：
      * ``计划购电量``：表头 + 334 天 × 144 个 10 分钟购电量 + ``全天购电量`` + ``全天购电费``；
      * ``充放电量``：334 天 × 6 个 4 小时块，日期只写在该块首行；``时刻``/``储电量`` 写在每日前两行
        （第 1 行 ``00:00:00`` 载 0:00 储电量、第 2 行 ``24:00`` 载 24:00 储电量）；
      * ``紧急购电量``：每日期至少 1 行；``J_d = 0`` 的日期写 ``时间段 = —``、``购电量 = 0``。

    所有写入都经 :func:`_assign`（含 ``None`` 也真正清空），并在保存前用 :func:`_residual_report`
    复核「应空却仍有值」的单元格；复核结果经返回值的 ``format_residuals`` 交给调用方做硬门禁。
    """
    days = len(dates)
    purchase = np.asarray(purchase_kwh, dtype=float)
    charge = np.asarray(charge_kwh, dtype=float)
    discharge = np.asarray(discharge_kwh, dtype=float)
    if purchase.shape != (days, PERIODS_PER_DAY):
        raise InputValidationError(f"purchase 形状 {purchase.shape} != ({days}, {PERIODS_PER_DAY})")
    if charge.shape != purchase.shape or discharge.shape != purchase.shape:
        raise InputValidationError("charge/discharge 形状与 purchase 不一致")
    if len(emergency_intervals) != days:
        raise InputValidationError("emergency_intervals 天数与 dates 不一致")
    if not template.is_file():
        raise InputValidationError(f"模板不存在：{template}")

    workbook = openpyxl.load_workbook(template)
    try:
        titles = list(workbook.sheetnames)
        for name in (PLAN_SHEET, STORAGE_SHEET, EMERGENCY_SHEET):
            if name not in titles:
                raise InputValidationError(f"模板缺少工作表『{name}』：{titles}")

        plan = workbook[PLAN_SHEET]
        if plan.max_row < days + 1:
            raise InputValidationError(f"模板『{PLAN_SHEET}』只有 {plan.max_row - 1} 行，少于 {days} 天")
        for index, date_text in enumerate(dates):
            row = 2 + index
            actual = _date_label(plan.cell(row=row, column=1).value)
            if actual != date_text:
                raise InputValidationError(f"模板『{PLAN_SHEET}』第 {row} 行日期 {actual!r} != {date_text!r}")
            for period in range(PERIODS_PER_DAY):
                _assign(plan, row=row, column=2 + period, value=float(purchase[index, period]))
            _assign(plan, row=row, column=PERIODS_PER_DAY + 2, value=float(day_energy_kwh[index]))
            _assign(plan, row=row, column=PERIODS_PER_DAY + 3, value=float(day_cost_yuan[index]))
        if plan.max_row > days + 1:
            _clear_rows(plan, first_row=days + 2, last_row=plan.max_row, last_column=PERIODS_PER_DAY + 3)

        storage = workbook[STORAGE_SHEET]
        storage_blocks = [str(storage.cell(row=2 + k, column=2).value) for k in range(BLOCKS_PER_DAY)]
        blocks_filled = 0
        for index, date_text in enumerate(dates):
            date_value = datetime.strptime(date_text, "%Y-%m-%d")
            for block in range(BLOCKS_PER_DAY):
                row = 2 + index * BLOCKS_PER_DAY + block
                start = block * BLOCK_ROWS
                stop = start + BLOCK_ROWS
                _assign(storage, row=row, column=1, value=date_value if block == 0 else None)
                _assign(storage, row=row, column=2, value=storage_blocks[block])
                _assign(storage, row=row, column=3, value=float(np.sum(charge[index, start:stop])))
                _assign(storage, row=row, column=4, value=float(np.sum(discharge[index, start:stop])))
                _assign(
                    storage,
                    row=row,
                    column=5,
                    value="00:00:00" if block == 0 else ("24:00" if block == 1 else None),
                )
                _assign(
                    storage,
                    row=row,
                    column=6,
                    value=float(state_start_kwh[index]) if block == 0 else (
                        float(state_end_kwh[index]) if block == 1 else None
                    ),
                )
                blocks_filled += 1
        expected_storage_rows = 1 + days * BLOCKS_PER_DAY
        if storage.max_row > expected_storage_rows:
            _clear_rows(storage, first_row=expected_storage_rows + 1, last_row=storage.max_row, last_column=6)

        emergency = workbook[EMERGENCY_SHEET]
        row = 2
        interval_rows = 0
        for index, date_text in enumerate(dates):
            date_value = datetime.strptime(date_text, "%Y-%m-%d")
            intervals = list(emergency_intervals[index])
            if not intervals:
                intervals = [(ZERO_INTERVAL_TEXT, 0.0)]
            for order, (slot, kwh) in enumerate(intervals):
                _assign(emergency, row=row, column=1, value=date_value if order == 0 else None)
                _assign(emergency, row=row, column=2, value=slot)
                _assign(emergency, row=row, column=3, value=float(kwh))
                row += 1
                interval_rows += 1
        expected_emergency_rows = row - 1
        if emergency.max_row > expected_emergency_rows:
            _clear_rows(emergency, first_row=expected_emergency_rows + 1, last_row=emergency.max_row, last_column=3)

        dest.parent.mkdir(parents=True, exist_ok=True)
        residuals = _residual_report(workbook=workbook, days=days, emergency_rows=interval_rows)
        workbook.save(dest)
    finally:
        workbook.close()
    return {
        "plan_rows_filled": days,
        "plan_periods_per_row": PERIODS_PER_DAY,
        "storage_blocks_filled": blocks_filled,
        "emergency_rows_filled": interval_rows,
        "storage_blocks": storage_blocks,
        "format_residuals": residuals,
    }


def write_json(path: Path, payload: Any) -> None:
    """写 JSON；显式拒绝 NaN/Inf，避免把非有限值伪装成合法结果。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False)
    path.write_text(text + "\n", encoding="utf-8")


def delivery_dates(all_dates: Sequence[str], start_day_index: int) -> list[str]:
    """交付期日期列表（默认 2025-02-01 … 2025-12-31，334 天）。"""
    return [str(item) for item in all_dates[start_day_index:]]


def day_offset(dates: Sequence[str], target: str) -> int:
    """在附件 2 的日期序列中定位目标的 0 基日索引；找不到时严格失败。"""
    for index, item in enumerate(dates):
        if item == target:
            return index
    raise InputValidationError(f"附件 2 日期序列中找不到 {target}")


__all__ = [
    "Attachment1",
    "Attachment2",
    "InputValidationError",
    "day_offset",
    "delivery_dates",
    "fill_result2_workbook",
    "inspect_template",
    "md5_of",
    "read_attachment1",
    "read_attachment2",
    "write_json",
]
